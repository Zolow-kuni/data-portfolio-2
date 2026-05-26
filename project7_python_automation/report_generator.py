"""
report_generator.py — Excel report builder for Project 7
Called by automation_pipeline.py
"""

import os
import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# ── Style constants ────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT   = Font(color="FFFFFF", bold=True, size=10)
ALT_FILL      = PatternFill("solid", fgColor="EBF3FB")
CURRENCY_FMT  = '#,##0.00'
PCT_FMT       = '0.00"%"'
THIN_BORDER   = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
TITLE_FONT    = Font(bold=True, size=13, color="1F4E79")
SECTION_FONT  = Font(bold=True, size=11, color="2E74B5")


def _style_header_row(ws, row: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_data_rows(ws, start_row: int, end_row: int, col_count: int, currency_cols=None):
    currency_cols = currency_cols or []
    for r in range(start_row, end_row + 1):
        fill = ALT_FILL if r % 2 == 0 else PatternFill()
        for col in range(1, col_count + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill      = fill
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if col in currency_cols:
                cell.number_format = CURRENCY_FMT


def _autofit(ws, min_width=10, max_width=45):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def _write_df(ws, df: pd.DataFrame, start_row: int, currency_cols: list = None):
    """Write a DataFrame starting at start_row, return the last row written."""
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=start_row, column=col_idx, value=header)
    _style_header_row(ws, start_row, len(headers))

    for r_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    currency_col_nums = []
    if currency_cols:
        for col_name in currency_cols:
            if col_name in headers:
                currency_col_nums.append(headers.index(col_name) + 1)

    _style_data_rows(ws, start_row + 1, start_row + len(df), len(headers), currency_col_nums)
    return start_row + len(df)


# ── Sheet builders ─────────────────────────────────────────────────────────────

def _sheet1_cleaned(wb: Workbook, df: pd.DataFrame):
    ws = wb.active
    ws.title = "Cleaned Data"
    ws.freeze_panes = "A2"

    # Drop internal flag columns from display
    display_cols = [c for c in df.columns if c not in ("calculation_error", "age_outlier", "quantity_error")]
    display_df = df[display_cols].copy()
    display_df["Date"] = display_df["Date"].dt.strftime("%Y-%m-%d")

    ws.cell(1, 1, "Retail Sales — Cleaned Dataset").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(display_cols))

    _write_df(ws, display_df, start_row=2, currency_cols=["Price per Unit", "Total Amount"])
    _autofit(ws)


def _sheet2_validation(wb: Workbook, report: dict):
    ws = wb.create_sheet("Validation Report")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18

    ws.cell(1, 1, "Data Quality Validation Report").font = TITLE_FONT
    ws.cell(2, 1, f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(italic=True, size=9, color="888888")

    rows = [
        ("Metric", "Value"),
        ("Total Rows Ingested", report["total_ingested"]),
    ]
    for d in report["dropped"]:
        rows.append((f"Rows Dropped — {d['reason']}", d["count"]))
    for f in report["flagged"]:
        rows.append((f"Rows Flagged — {f['reason']}", f["count"]))
    rows.append(("Final Clean Row Count", report["final_clean_count"]))
    rows.append(("Data Quality Score (%)", report["data_quality_score"]))

    start = 4
    for r_idx, (label, value) in enumerate(rows, start):
        ws.cell(r_idx, 1, label)
        ws.cell(r_idx, 2, value)
        if r_idx == start:  # header
            ws.cell(r_idx, 1).font = HEADER_FONT
            ws.cell(r_idx, 1).fill = HEADER_FILL
            ws.cell(r_idx, 2).font = HEADER_FONT
            ws.cell(r_idx, 2).fill = HEADER_FILL
        elif r_idx % 2 == 0:
            ws.cell(r_idx, 1).fill = ALT_FILL
            ws.cell(r_idx, 2).fill = ALT_FILL

    # Highlight quality score cell
    score_row = start + len(rows) - 1
    score_cell = ws.cell(score_row, 2)
    score = report["data_quality_score"]
    score_cell.font = Font(bold=True, color="FFFFFF")
    score_cell.fill = PatternFill("solid", fgColor="1E8449" if score >= 95 else "F39C12" if score >= 80 else "C0392B")


def _sheet3_category(wb: Workbook, rev_cat: pd.DataFrame):
    ws = wb.create_sheet("Revenue by Category")
    ws.cell(1, 1, "Revenue by Product Category").font = TITLE_FONT

    _write_df(ws, rev_cat, start_row=3, currency_cols=["Revenue"])
    _autofit(ws)


def _sheet4_monthly(wb: Workbook, monthly: pd.DataFrame):
    ws = wb.create_sheet("Monthly Trend")
    ws.cell(1, 1, "Monthly Revenue Trend with MoM Growth").font = TITLE_FONT

    _write_df(ws, monthly, start_row=3, currency_cols=["Revenue"])
    _autofit(ws)


def _sheet5_customers(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Customer Summary")
    ws.cell(1, 1, "Customer Analysis").font = TITLE_FONT

    row = 3
    ws.cell(row, 1, "Top 10 Customers by Total Spend").font = SECTION_FONT
    row += 1
    last = _write_df(ws, analysis["top_customers"], start_row=row, currency_cols=["Total Spend"])

    row = last + 3
    ws.cell(row, 1, "Age Group Spending Patterns").font = SECTION_FONT
    row += 1
    last = _write_df(ws, analysis["age_group_spending"], start_row=row, currency_cols=["Revenue", "Avg_Order"])

    row = last + 3
    ws.cell(row, 1, "Gender × Product Category Revenue").font = SECTION_FONT
    row += 1
    _write_df(ws, analysis["gender_by_category"], start_row=row, currency_cols=["Revenue"])

    _autofit(ws)


# ── Main entry point ───────────────────────────────────────────────────────────

def generate_excel_report(df: pd.DataFrame, validation_report: dict, analysis: dict, output_dir: str) -> str:
    today = datetime.date.today().strftime("%Y-%m-%d")
    filename = f"Retail_Sales_Report_{today}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()

    _sheet1_cleaned(wb, df)
    _sheet2_validation(wb, validation_report)
    _sheet3_category(wb, analysis["revenue_by_category"])
    _sheet4_monthly(wb, analysis["monthly_revenue"])
    _sheet5_customers(wb, analysis)

    wb.save(filepath)
    return filepath
